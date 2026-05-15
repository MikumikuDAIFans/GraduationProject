"""
Fix run-level formatting for all heading paragraphs.
The style definitions are correct, but runs have explicit font overrides
that need to be fixed: ascii should be Times New Roman (not 黑体),
and bold should be set for Heading1/2.
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = 'docs/academic/PROPOSAL/论文修改/毕业设计-张嘉洋_项目现状对齐修订版.docx'
doc = Document(DOC_PATH)

# Style -> font spec
heading_specs = {
    'ThesisHeading1': {'ascii': 'Times New Roman', 'ea': '黑体', 'cs': 'Times New Roman', 'sz': 30, 'bold': True},
    'ThesisHeading2': {'ascii': 'Times New Roman', 'ea': '黑体', 'cs': 'Times New Roman', 'sz': 28, 'bold': True},
    'ThesisHeading3': {'ascii': 'Times New Roman', 'ea': '黑体', 'cs': 'Times New Roman', 'sz': 24, 'bold': False},
}

fixed = 0
for p in doc.paragraphs:
    style_name = p.style.name
    if style_name not in heading_specs:
        continue

    spec = heading_specs[style_name]

    for run in p.runs:
        rPr = run._element.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            run._element.insert(0, rPr)

        # Fix fonts
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:ascii'), spec['ascii'])
        rFonts.set(qn('w:hAnsi'), spec['ascii'])
        rFonts.set(qn('w:eastAsia'), spec['ea'])
        rFonts.set(qn('w:cs'), spec['cs'])

        # Fix size
        sz = rPr.find(qn('w:sz'))
        if sz is None:
            sz = OxmlElement('w:sz')
            rPr.append(sz)
        sz.set(qn('w:val'), str(spec['sz']))

        szCs = rPr.find(qn('w:szCs'))
        if szCs is None:
            szCs = OxmlElement('w:szCs')
            rPr.append(szCs)
        szCs.set(qn('w:val'), str(spec['sz']))

        # Fix bold
        b = rPr.find(qn('w:b'))
        if spec['bold']:
            if b is None:
                b = OxmlElement('w:b')
                rPr.append(b)
        else:
            if b is not None:
                rPr.remove(b)

        fixed += 1

doc.save(DOC_PATH)
print(f'Fixed {fixed} heading runs across all ThesisHeading1/2/3 paragraphs')

# Verify
doc2 = Document(DOC_PATH)
for sn in ['ThesisHeading1', 'ThesisHeading2', 'ThesisHeading3']:
    for p in doc2.paragraphs:
        if p.style.name == sn and p.text.strip():
            rPr = p.runs[0]._element.find(qn('w:rPr')) if p.runs else None
            if rPr is not None:
                rFonts = rPr.find(qn('w:rFonts'))
                sz = rPr.find(qn('w:sz'))
                b = rPr.find(qn('w:b'))
                ascii_font = rFonts.get(qn('w:ascii')) if rFonts is not None else '?'
                ea_font = rFonts.get(qn('w:eastAsia')) if rFonts is not None else '?'
                size = sz.get(qn('w:val')) if sz is not None else '?'
                bold = b is not None
                print(f'  {sn}: "{p.text[:30]}" | ascii={ascii_font} ea={ea_font} size={size} bold={bold}')
            break
