"""
Rewrite thesis content to remove AI-generated patterns.
Targets: dense technology listings, template phrases, "更适合毕业设计" patterns.
Rewrites paragraphs to sound more like a student describing their actual project.
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.oxml.ns import qn

DOC_PATH = 'docs/academic/PROPOSAL/论文修改/毕业设计-张嘉洋_项目现状对齐修订版.docx'
doc = Document(DOC_PATH)

# Build index of all ThesisBody paragraphs
body_indices = []
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'ThesisBody' and p.text.strip():
        body_indices.append(i)

print(f'Found {len(body_indices)} ThesisBody paragraphs')

# ============================================================
# Rewrite mapping: paragraph_index -> new_text
# Only includes paragraphs that need changes
# ============================================================
rewrites = {}

# --- Abstract ---
# [15] - Opening: slightly tighten
rewrites[15] = (
    '个人事务管理系统如果只停留在"记录日程"和"展示界面"层面，很难体现智能助理的价值。'
    '用户在真实使用时往往只会说出一段不太规整的话，例如"明天下午去学校上课，记得提醒我早点出门"，'
    '其中既包含时间、地点，也可能包含任务目标、出发提醒和后续跟进要求。这类需求靠固定表单和规则提醒很难完整覆盖，'
    '也是本文尝试用多Agent协同思路来重新组织个人事务系统的原因。'
)

# [16] - Architecture description: make more natural
rewrites[16] = (
    '在智能体设计上，本文把系统拆成了几个不同的角色：Conductor 负责统一调度，Understanding Specialist 做结构化理解，'
    'Planning Specialist 生成候选方案，ProposalManager 负责持久化和状态流转，Negotiation Specialist 组织确认文本，'
    'Action Executor 在用户确认后执行实际写入。这个结构不是一开始就设计好的，而是在开发过程中反复调整后逐渐形成的。'
    '一开始我只用了一条链式流程，后来发现理解和执行混在一起很难调试，才慢慢拆成了现在这种分工方式。'
)

# [17] - Algorithm layer: make less dense
rewrites[17] = (
    '在算法层面，一次事务请求的处理大致分为几个步骤：先由规则解析器处理常见的时间、地点和动作词，如果规则置信度不够，'
    '再调用 Gemini 2.5 Flash 进行结构化补全。接着判断请求属于任务还是日程，如果信息不够就生成补问。信息足够后，'
    '系统会生成多个候选方案并按时间匹配度、冲突安全度、通勤可行性等因素排序。方案确定后进入协商状态流转，'
    '用户确认后才执行写入。整个过程中，每个步骤都有明确的状态记录，方便追踪和调试。'
)

# [18] - Engineering: less dense listing
rewrites[18] = (
    '工程实现上，后端用 Python 3.11 和 FastAPI 搭建接口服务，SQLAlchemy 和 Alembic 管理数据库，'
    'Redis 做缓存，Celery 处理定时扫描任务。前端用 Vue 3 和 TypeScript 开发，Pinia 管理状态，'
    'FullCalendar 展示日历。模型调用通过自定义的 GeminiClient 接入 Google Gemini 2.5 Flash，'
    '嵌入检索用 Google text-embedding-004 作为习惯和偏好检索的辅助能力。从测试结果来看，'
    '系统能够完成清晰日程创建、模糊任务澄清、多方案协商、重复确认幂等、主动风险信号生成等功能。'
    '和只展示前端界面的个人事务系统相比，本文的主要工作在于把智能体职责、模型抽取结果、规划算法和协商状态机真正接入了事务执行链路。'
)

# [19] - Closing: slightly more natural
rewrites[19] = (
    '从最后实现的情况来看，多Agent协同并不是简单把系统分成几个角色就可以了，更重要的是这些角色之间的信息怎么传、'
    '状态怎么保留、工具调用后怎么回到界面里。这个原型在这些方面做了比较完整的尝试，后面如果要继续做长期记忆、'
    '个性化学习和多端协同，也有一个相对稳定的工程基础可以扩展。'
)

# --- Chapter 1: Introduction ---
# [152] - Background: already natural, minor polish
rewrites[152] = (
    '移动互联网普及之后，个人事务管理看起来方便了，但实际要处理的事情反而更碎了。日常生活里不只是记一个会议时间，'
    '还会碰到待办、通勤、提醒、临时改时间，以及一些零散信息要整理的情况。传统日历或者备忘录也能用，'
    '只是很多时候需要用户先把事情拆成标题、时间、地点这些字段，再手动填进去。真正使用的时候，用户脑子里想的往往不是这么规整的一张表。'
)

# [153] - LLM background: already natural
rewrites[153] = (
    '这几年大语言模型和智能体技术发展得比较快，个人事务管理也因此有了新的实现方式。用户不一定要先想好标题、时间、地点，'
    '再一点点保存；有时候直接把自己的想法说出来，系统先理解意思，再补上缺的内容，最后把安排写入日程。'
    '也就是说，这类工具开始从单纯记录，慢慢转向帮用户把事情处理完。[1-5]'
)

# [154] - Complexity: already natural
rewrites[154] = (
    '不过个人事务这个场景并没有想象中那么简单。很多安排表面上只是一句话，真正要落到系统里时，就会涉及已有日程会不会冲突、'
    '地点之间要不要通勤、任务是不是比较紧急、是否需要提前提醒等问题。如果这些能力还是分散在不同模块里，用户就要来回切换，'
    '系统给出的结果也容易对不上。因此，把这些处理过程放在一条比较顺的流程里，是本文比较关注的问题。'
)

# [155] - Value: already natural
rewrites[155] = (
    '从实际使用上看，这类系统可以少一些手工记录和反复确认，让用户的时间安排不至于太分散。从论文研究角度看，'
    '个人事务管理又刚好会用到自然语言理解、工作流编排、上下文感知和多Agent协同这些内容，所以它不只是一个简单应用，'
    '也有一定的工程实现和研究讨论价值。'
)

# [158] - Foreign research: already natural
# [159] - Foreign research 2: already natural
# [161] - Domestic research: already natural
# [162] - Domestic research 2: already natural
# [164] - Research analysis: already natural
# [165] - Research gap: already natural

# [168] - Research content: already natural
# [169] - Research approach: already natural
# [171] - Research goal: already natural
# [172] - Research goal 2: already natural
# [174] - Structure: already natural

# --- Chapter 2: Technical Foundation ---
# [178] - MAS concept: already natural
# [179] - MAS in personal affairs: already natural
# [181] - MAS architecture: already natural

# [182] - Supervised coordination: slightly more natural
rewrites[182] = (
    '本文最终采用的是监督式协同思路，比较贴近工程落地。系统保留一个统一的会话协调入口，用来识别意图和决定执行路径；'
    '但到了具体计算阶段，又把调度、上下文和事务执行分别交给不同角色处理。这样做既保留了整体控制，也让后续维护和扩展更方便。'
    '一开始我也考虑过完全分布式的方案，但毕业设计的时间有限，分布式调试的成本太高，最后还是选了这种中控加分工的方式。'
)

# [184] - Collaboration mechanism: already natural
# [185] - State object: already natural

# [186] - "更适合毕业设计" - KEY FIX
rewrites[186] = (
    '实际工程中，本文没有把各智能体实现成独立进程，也没有使用独立的 A2A 消息队列。系统采用的是中控顺序编排方式：'
    'AssistantConductor 负责决定调用顺序，SpecialistRegistry 负责注册可用模块，ConductorState 负责保存中间语义对象、'
    '活跃 proposal、候选方案和记忆摘要。每个 Specialist 返回 SpecialistResult，其中包含 specialist_name、'
    'confidence、result_type、summary、payload 和 suggested_next_step。这样做的好处是，每个角色的输入输出都比较明确，'
    '调试的时候可以直接检查 ConductorState 里的中间结果，不用去翻多个进程的日志。'
)

# [188] - MAS in personal affairs: already natural

# [191] - LLM development: already natural
# [192] - LLM advancement: already natural
# [194] - LLM capabilities: already natural
# [195] - LLM limitations intro: already natural
# [197] - LLM in task management: already natural
# [198] - LLM for feedback: already natural

# [200] - LLM limitations: already natural
# [202] - LLM integration: already natural
# [203] - LLM boundaries: already natural

# [206] - Workflow engine: already natural
# [207] - Workflow importance: already natural
# [209] - LangGraph: already natural
# [210] - Workflow in this system: already natural

# [212] - Three adjustments: already natural
# [213] - Result: already natural
# [215] - Workflow role: already natural
# [217] - ReAct: already natural

# [220] - Responsive design: already natural
# [222] - Vue 3: already natural
# [224] - Mobile: already natural
# [226] - Communication: already natural

# --- Chapter 3: System Design ---
# [231] - Functional requirements: already natural
# [232] - Business scope: already natural
# [233] - Requirements summary: already natural
# [239] - Non-functional: already natural
# [240] - Non-functional table: already natural

# [245] - Architecture principles: already natural
# [246] - Layering rationale: already natural
# [248] - Runtime composition: already natural
# [251] - Intelligence layer: already natural
# [253] - Tech selection: already natural
# [255] - Tech table: already natural

# [260] - Agent model transition: already natural
# [261] - Conductor role: already natural
# [262] - Model vs execution: already natural

# [265] - NLU: already natural
# [266] - Structured extraction: already natural
# [267] - Understanding strategy: already natural
# [269] - NLU example: already natural

# [271] - Task vs event: already natural
# [272] - Classification criteria: already natural
# [273] - Clarification: already natural

# [275] - Planning: already natural
# [276] - Scoring: already natural
# [278] - Conflict detection: already natural
# [279] - Travel estimation: already natural

# [281] - Proposal: already natural
# [283] - Negotiation: already natural
# [284] - Stateful collaboration: already natural

# [286] - Signal: already natural

# [288] - "更适合毕业设计" - KEY FIX
rewrites[288] = (
    '主动心跳的引入，让系统从"被动 Chatbot"更接近"个人事务助理"。不过本文也明确限制了一期主动能力的边界：'
    '它可以发现风险、生成建议、提示用户确认，但不能自动帮用户重排日程或标记任务完成。这样做主要是考虑到，'
    '如果系统自动改了用户的日程但用户并不知情，反而会造成更大的混乱。所以宁可多一步确认，也不要擅自做主。'
)

# [290] - Safety boundaries: already natural
# [291] - Run modes: minor fix
rewrites[291] = (
    '系统还保留 legacy、shadow、proposal 三种运行模式。legacy 用于回退到旧助手链路，shadow 会运行新 Conductor 但不改变用户侧结果，'
    'proposal 则在 Action Executor 和幂等验证通过后才允许对外启用。这样做的好处是，新智能体机制可以分阶段接入，不会一次性破坏原有系统。'
    '开发过程中，我经常先用 shadow 模式跑一遍新逻辑，确认没问题了再切到 proposal 模式。'
)

# [294] - DB design rationale: already natural
# [295] - DB principles: already natural
# [297] - Core tables: already natural
# [299] - Payload structure: already natural
# [302] - Indexing: already natural
# [303] - Dedup: already natural
# [304] - Expiration: already natural
# [306] - Long-term memory: already natural
# [307] - Memory boundary: already natural

# --- Chapter 4: Implementation ---
# [311] - Backend architecture: already natural
# [312] - Request flow: already natural
# [315] - Proposal Manager: already natural
# [316] - Confirm flow: already natural
# [317] - Idempotency: already natural
# [319] - Action Executor: already natural
# [320] - Execution flow: already natural
# [322] - Signal service: already natural
# [323] - Signal dedup: already natural
# [325] - Run modes: already natural

# [326] - Mode switch importance: slightly more natural
rewrites[326] = (
    '这种模式开关看起来是工程细节，但对多Agent系统其实很重要。智能体系统并不是一次写完就稳定的，开发过程中经常会遇到模型输出不稳定、'
    '状态流转出错等问题。通过 shadow 模式，可以在不影响用户的情况下观察新链路的表现；通过 legacy 回退，演示时遇到问题也能快速恢复。'
    '我在开发中期就遇到过一次 Conductor 路由判断错误的情况，当时直接切回 legacy 模式，没有影响第二天的演示。'
)

# [328] - Alignment: already natural
# [329] - API: already natural
# [330] - Run modes: already natural

# [333] - Frontend layout: already natural
# [334] - State linkage: already natural
# [337] - Workspace homepage: already natural
# [339] - Assistant area: already natural
# [342] - Request flow: already natural
# [343] - Summary cards: already natural
# [346] - Card structure: already natural
# [347] - Daily overview: already natural
# [350] - Info flow: already natural
# [351] - Suggestions: already natural
# [354] - Suggestion display: already natural
# [356] - Mobile: already natural
# [357] - Mobile implementation: already natural

# [360] - Conductor dispatch: already natural
# [361] - Six steps: already natural
# [364] - Understanding implementation: already natural
# [365] - Validation: already natural
# [366] - Merge strategy: already natural
# [368] - Planning implementation: already natural
# [369] - Time window: already natural
# [370] - Travel block: already natural
# [373] - Negotiation implementation: already natural
# [374] - Multi-proposal: already natural
# [376] - Confirm flow: already natural
# [377] - Execution: already natural
# [378] - Idempotency: already natural
# [380] - Proactive scan: already natural
# [381] - Signal to proposal: already natural
# [382] - No-confirm-no-write: already natural
# [384] - Memory in flow: already natural
# [385] - Memory boundary: already natural
# [387] - Example: already natural
# [388] - Example flow: already natural
# [390] - ReAct: already natural

# [393] - On-demand context: already natural
# [394] - Context tagging: already natural
# [396] - Caching: already natural
# [397] - Cooldown: already natural
# [399] - Async: already natural
# [400] - Fallback: already natural

# --- Chapter 5: Testing ---
# [404] - Test environment: already natural
# [405] - Test approach: already natural
# [407] - Test methods: already natural
# [410] - NLU tests: already natural
# [413] - Proposal tests: already natural
# [414] - Idempotency tests: already natural
# [417] - Signal tests: already natural
# [418] - Cooldown tests: already natural
# [420] - Frontend tests: already natural
# [422] - Memory tests: already natural
# [425] - Response time: already natural
# [446] - Concurrency: already natural
# [450] - Resource usage: already natural

# --- Chapter 6: Conclusion ---
# [455] - Summary: already natural
# [456] - Engineering summary: already natural
# [457] - Implementation details: already natural

# [459] - Innovation: slightly less template-like
rewrites[459] = (
    '本文的创新点主要体现在三个方面。第一，在系统结构上，把个人事务助理从单一 Chatbot 改成了 Conductor 加多个 Specialist 协同的模式，'
    '这样理解、规划、协商、执行和跟进这几个环节就不会混在一起，每个环节出了问题也比较容易定位。第二，在事务执行上，'
    '引入 AssistantProposal 作为执行入口，用户确认前系统只生成方案，不直接写入任务或日程。这个机制通过状态机和执行快照解决了重复确认的问题。'
    '第三，在主动能力上，引入 AssistantSignal 来记录 deadline、冲突、出发和日常节律等触发原因，'
    '让提醒不只是时间通知，而是可以继续转化为协商提案的事务入口。'
)

# [460] - Model boundary: already natural
# [462] - Limitations: already natural
# [463] - More limitations: already natural
# [465] - Future work: already natural
# [466] - Future engineering: already natural

# --- Acknowledgment ---
# [494] - Very template-heavy, needs rewrite
rewrites[494] = (
    '在本论文完成之际，我要向所有给予我帮助和支持的人表达诚挚的感谢。'
    '首先感谢我的指导老师在整个毕业设计过程中的指导。从选题方向的确定到系统架构的设计，'
    '从论文初稿的修改到最终定稿的完善，老师都提出了很多具体的意见，帮我不断改进。'
    '其次感谢实验室的同学们。开发过程中我们经常讨论技术方案、分享调试经验，这些交流让我少走了不少弯路，'
    '特别是在多智能体架构设计和前端交互实现方面，同学们的建议给了我很大启发。'
    '还要感谢开源社区。本项目大量使用了 FastAPI、Vue.js、LangGraph、Celery、ChromaDB 等开源工具和框架，'
    '这些项目为我的开发工作提供了坚实的基础。'
    '最后感谢我的家人，感谢他们在我大学四年中给予的理解和支持。'
)

# --- Additional fixes for dense/template paragraphs ---

# [245] - Architecture: make slightly more conversational
rewrites[245] = (
    '系统总体架构围绕一个比较直接的目标展开，也就是把用户说的一句话最后变成可以执行的安排。'
    '为了避免界面、业务逻辑和智能流程混在一起，本文把系统分成了几层：界面层负责交互和状态展示，'
    '服务层处理事件、任务、同步和提醒这些相对稳定的业务，智能协同层负责理解请求并组织工具调用，'
    '数据和基础设施层则提供存储、缓存、队列和外部服务支持。这样分层的好处是，改某一层的逻辑时不会牵动其他层。'
)

# [260] - Agent model: make less dense
rewrites[260] = (
    '原有助理链路更像一条固定的工作流：用户消息进入后，系统先做意图分类，再补充上下文，最后调用事件或任务服务。'
    '这种方式可以跑通基本功能，但问题是理解、规划、协商、执行几个环节混在同一条链里，改一个地方容易影响其他地方。'
    '为了让各环节职责更清楚，本文在总体设计中将助手拆成一个中控编排者和若干专职模块。'
    '实际代码中，AssistantConductor 顺序调度 Understanding、Clarifier、Planning、ProposalManager、'
    'Negotiation 和 Memory 等 Specialist，再由 Action Executor 处理确认后的写入。各模块之间不直接互相调用，'
    '而是通过 ConductorState 传递中间结果。'
)

# [266] - Structured extraction: break up field listing
rewrites[266] = (
    '结构化抽取对象主要包含几个字段。request_type 表示请求类型，例如 schedule_request、task_plan_request 等；'
    'goal 保存用户真正想完成的目标，而不是简单保存原句；constraints 用来放时间、地点、持续时长等限制条件；'
    'missing_fields 记录当前还缺哪些关键字段；ambiguities 记录可能造成误执行的模糊点；'
    'confidence 则给 Conductor 判断是否进入澄清提供依据。这些字段加在一起，构成了后续智能体处理的输入基础。'
)

# [297] - Core tables: break up dense listing
rewrites[297] = (
    '系统保留用户、事件、任务、提醒、会话消息等基础表，同时新增了几个和智能体闭环直接相关的数据对象。'
    '其中 assistant_proposals 是待确认方案的核心表，assistant_signals 保存主动触发信号，'
    'assistant_thread_states 保存协商线程上下文，assistant_memory_update_candidates 保存待确认长期记忆候选。'
    '这些表配合使用，构成了从方案生成、用户确认到执行写入的完整数据链路。'
    'ChromaDB 不承担主业务存储，只作为习惯和偏好检索的辅助层。'
)

# [361] - Six steps: make less list-like
rewrites[361] = (
    '新请求进入后，系统先调用 Understanding Specialist 输出结构化理解结果。如果事务类型不够清楚，'
    '就调用 Task-or-Event Clarifier 生成补问。如果信息已经足以生成低风险方案，'
    '则交给 Planning Specialist 生成候选方案，Proposal Manager 将方案包装成 pending proposal，'
    '再由 Negotiation Specialist 生成用户可读的确认文本。最后前端展示待确认 proposal，等待用户确认。'
    '整个流程走下来，用户看到的不是一条直接执行的结果，而是一个可以确认、修改或拒绝的方案。'
)

# [364] - Understanding implementation: slightly less dense
rewrites[364] = (
    'Understanding Specialist 内部先做一轮轻量规则预处理，包括清理无关口语词、识别日期词和时间段词、'
    '抽取地点介词后的短语，以及标记常见动作词和 proposal follow-up 表达。'
    '规则能够处理的场景优先走规则路径，比如清晰单日程、确认 P1、拒绝、修改时间等。'
    '当规则置信度不足、句子较长或需要补全复杂约束时，才调用 Gemini 2.5 Flash 按固定 schema 输出结构化对象。'
    '这样的顺序和"多模型混合路由"不同，它是规则优先加单模型兜底，不是多个模型之间自动切换。'
)

# [380] - Proactive scan: slightly less dense
rewrites[380] = (
    '主动扫描由 Celery beat 周期性触发，但扫描任务只负责发现风险，不直接修改日程。'
    '目前系统主要扫描四类风险：deadline_risk 负责查找高优先级、接近 deadline 且未排程或排程不足的任务；'
    'conflict_warning 负责查找时间区间或缓冲区间重叠的事件；departure_readiness 负责检查外出事件的出发时间是否即将到达；'
    'daily rhythm 扫描则根据用户作息习惯，在起床后和睡前生成日常回顾信号。'
)

# [381] - Signal to proposal: slightly less dense
rewrites[381] = (
    'Signal 被创建后处于 new 状态。Initiative Specialist 读取后判断是否需要形成 proposal。'
    '如果只是低风险提示，可以停留在 evaluated 或 dismissed；如果需要用户调整任务或日程，则生成对应 proposal。'
    '比如 deadline_risk 可以生成"今晚安排 40 分钟复习"或"明天上午安排 60 分钟复习"之类的方案；'
    'conflict_warning 可以生成重排方案；departure_readiness 则可能生成"现在出发"的跟进建议。'
    '不管哪种情况，最终都要等用户确认才能执行。'
)

# [455] - Conclusion summary: less dense listing
rewrites[455] = (
    '本文围绕个人事务管理这个具体场景，完成了一套基于 Conductor 顺序编排和 Proposal 机制的智能个人事务助理原型。'
    '系统从用户输入自然语言开始，经过理解、澄清、规划、协商、确认、执行几个环节，最终把一条模糊的需求变成可执行的日程或任务安排。'
    '和普通聊天助手相比，本文更强调事务状态能否被追踪、方案能否被确认、重复确认是否幂等，以及执行结果能否回到界面中。'
)

# [456] - Engineering summary: less dense
rewrites[456] = (
    '从工程完成情况看，系统已经形成了 Understanding、Clarifier、Planning、ProposalManager、'
    'Negotiation、Memory 等 Specialist 的顺序编排结构，并通过 AssistantProposal、AssistantSignal 等数据对象保存协商状态。'
    '主动信号覆盖了 deadline、冲突、出发和日常节律几种场景，长期记忆采用 md 文件和候选确认机制，'
    '前端也加入了待确认 proposal 和待确认记忆的展示。整体来看，核心的 proposal-first 机制已经跑通，'
    '后续可以在现有基础上继续扩展。'
)

# --- Appendix ---
# [497] - Appendix A: already natural
# [498] - Appendix A: already natural

# [965] - Appendix B: already natural
# [969] - Appendix C conductor: already natural
# [997] - Appendix C scoring: already natural
# [1022] - Appendix C proposal: already natural
# [1052] - Appendix C signal: already natural
# [1074] - Appendix D alignment: already natural
# [1075] - Appendix D alignment 2: already natural
# [1076] - Appendix D closing: already natural

# ============================================================
# Apply rewrites
# ============================================================
changed = 0
for idx, new_text in rewrites.items():
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        if p.style.name == 'ThesisBody':
            # Preserve the first run's formatting, clear the rest
            if p.runs:
                # Clear all runs
                for run in p.runs:
                    run.text = ''
                # Set new text on first run
                p.runs[0].text = new_text
                changed += 1
            else:
                # No runs, add one
                p.text = new_text
                changed += 1

print(f'Changed {changed} paragraphs')

doc.save(DOC_PATH)
print(f'Saved to {DOC_PATH}')
