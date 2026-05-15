import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document('docs/academic/PROPOSAL/论文修改/毕业设计-张嘉洋_项目现状对齐修订版.docx')

def fix_table_borders(table):
    """Fix borders to match template: top/bottom=18, insideH=4, no vertical."""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Remove existing borders
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)

    borders = OxmlElement('w:tblBorders')

    # Top: lineWidth=18 (2.25pt thick line)
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '18')
    top.set(qn('w:space'), '0')
    top.set(qn('w:color'), '000000')
    borders.append(top)

    # Bottom: lineWidth=18
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), '000000')
    borders.append(bottom)

    # Inside horizontal: lineWidth=4 (thin line between rows)
    insideH = OxmlElement('w:insideH')
    insideH.set(qn('w:val'), 'single')
    insideH.set(qn('w:sz'), '4')
    insideH.set(qn('w:space'), '0')
    insideH.set(qn('w:color'), '000000')
    borders.append(insideH)

    # No vertical borders
    for side in ['left', 'right', 'insideV']:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)

    tblPr.append(borders)

def remove_table_style(table):
    """Remove aff1 or other table styles."""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is not None:
        tblStyle = tblPr.find(qn('w:tblStyle'))
        if tblStyle is not None:
            tblPr.remove(tblStyle)
        # Also remove tblLook which references style banding
        tblLook = tblPr.find(qn('w:tblLook'))
        if tblLook is not None:
            tblPr.remove(tblLook)

def fix_cell_format(cell, is_header=False):
    """Fix cell font and alignment to match template."""
    for p in cell.paragraphs:
        # Center align all cells (matching template)
        p.alignment = 1  # CENTER

        for run in p.runs:
            # Set font size to 10.5pt (5号)
            run.font.size = Pt(10.5)

            # Header: bold. Data: not bold
            run.font.bold = True if is_header else None

            # Set font name based on content
            text = run.text
            has_chinese = any('\u4e00' <= c <= '\u9fff' for c in text)
            if has_chinese:
                font_name = '宋体'
            else:
                font_name = 'Times New Roman'

            run.font.name = font_name
            rPr = run._element.find(qn('w:rPr'))
            if rPr is not None:
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is None:
                    rFonts = OxmlElement('w:rFonts')
                    rPr.append(rFonts)
                rFonts.set(qn('w:ascii'), font_name)
                rFonts.set(qn('w:hAnsi'), font_name)
                rFonts.set(qn('w:eastAsia'), '宋体')

# Apply fixes to all 19 tables
for t_idx, table in enumerate(doc.tables):
    # 1. Remove table style (aff1 etc.)
    remove_table_style(table)

    # 2. Fix borders to template format
    fix_table_borders(table)

    # 3. Fix header row
    for cell in table.rows[0].cells:
        fix_cell_format(cell, is_header=True)

    # 4. Fix data rows
    for row in table.rows[1:]:
        for cell in row.cells:
            fix_cell_format(cell, is_header=False)

    print(f'Table[{t_idx}]: fixed')

doc.save('docs/academic/PROPOSAL/论文修改/毕业设计-张嘉洋_项目现状对齐修订版.docx')
print('\nAll 19 tables fixed successfully')
